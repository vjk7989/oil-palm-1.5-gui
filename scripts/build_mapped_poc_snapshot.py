"""Build the portable Mapped POC snapshot from read-only survey sources."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook
from PIL import Image, ImageOps


SOURCE_AREA_COUNTS = [27, 13, 17, 53, 2]
DISPLAY_AREA_COUNTS = [52, 13, 17, 53, 2]
WORKBOOK_SHA256 = "787D536D77E3883B3DD88C0B22A8C2144EEF4AAA3903FF15A5236D9BF8891EED"
EVIDENCE_MAX_EDGE = 1280
EVIDENCE_QUALITY = 76
GEOFENCE_PAD_DEGREES = 0.000015


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def workbook_rows(path: Path) -> list[dict[str, object]]:
    if sha256(path) != WORKBOOK_SHA256:
        raise ValueError("Workbook checksum does not match the approved source")
    sheet = load_workbook(path, data_only=True, read_only=True)["Tree Health"]
    headers = [sheet.cell(4, column).value for column in range(1, 25)]
    rows = [
        dict(zip(headers, values))
        for values in sheet.iter_rows(min_row=5, max_row=116, max_col=24, values_only=True)
    ]
    if [row["Tree ID"] for row in rows] != [f"TREE-{index:04d}" for index in range(1, 113)]:
        raise ValueError("Expected the approved TREE-0001 through TREE-0112 sequence")
    return rows


def inventory_rows(path: Path) -> list[dict[str, object]]:
    rows = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(rows, list):
        raise ValueError("GPS inventory must contain a JSON array")
    return rows


def unique_green_captures(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    green = [
        row
        for row in rows
        if row.get("image_role") == "MS_G"
        and row.get("capture_uuid")
        and row.get("latitude") is not None
        and row.get("longitude") is not None
    ]
    green.sort(key=lambda row: (row["folder_index"], row["capture_number"], row["capture_timestamp"], row["file_name"]))
    unique: list[dict[str, object]] = []
    seen: set[str] = set()
    for row in green:
        capture_uuid = str(row["capture_uuid"])
        if capture_uuid in seen:
            continue
        seen.add(capture_uuid)
        unique.append(row)
    counts = [sum(row["folder_index"] == index for row in unique) for index in range(1, 6)]
    if len(unique) != 112 or counts != SOURCE_AREA_COUNTS:
        raise ValueError(f"Expected 112 unique captures split {SOURCE_AREA_COUNTS}; received {counts}")
    return unique


def convex_hull(points: list[tuple[float, float]]) -> list[dict[str, float]]:
    """Return a stable display footprint around the camera-position point set."""
    expanded = {
        (longitude + dx, latitude + dy)
        for latitude, longitude in points
        for dx, dy in (
            (-GEOFENCE_PAD_DEGREES, -GEOFENCE_PAD_DEGREES),
            (-GEOFENCE_PAD_DEGREES, GEOFENCE_PAD_DEGREES),
            (GEOFENCE_PAD_DEGREES, -GEOFENCE_PAD_DEGREES),
            (GEOFENCE_PAD_DEGREES, GEOFENCE_PAD_DEGREES),
        )
    }
    ordered = sorted(expanded)

    def cross(origin: tuple[float, float], a: tuple[float, float], b: tuple[float, float]) -> float:
        return (a[0] - origin[0]) * (b[1] - origin[1]) - (a[1] - origin[1]) * (b[0] - origin[0])

    lower: list[tuple[float, float]] = []
    for point in ordered:
        while len(lower) >= 2 and cross(lower[-2], lower[-1], point) <= 0:
            lower.pop()
        lower.append(point)
    upper: list[tuple[float, float]] = []
    for point in reversed(ordered):
        while len(upper) >= 2 and cross(upper[-2], upper[-1], point) <= 0:
            upper.pop()
        upper.append(point)
    hull = lower[:-1] + upper[:-1]
    return [{"latitude": round(latitude, 9), "longitude": round(longitude, 9)} for longitude, latitude in hull]


def evidence_source_for(capture: dict[str, object], rows: list[dict[str, object]], mission_root: Path) -> Path:
    candidates = [
        row
        for row in rows
        if row.get("image_role") == "D"
        and row.get("folder_index") == capture["folder_index"]
        and row.get("capture_number") == capture["capture_number"]
    ]
    if len(candidates) != 1:
        raise ValueError(
            f"Expected one natural-colour image for folder {capture['folder_index']} "
            f"capture {capture['capture_number']}; received {len(candidates)}"
        )
    source = mission_root / str(candidates[0]["folder"]) / str(candidates[0]["file_name"])
    if not source.is_file():
        raise FileNotFoundError(f"Natural-colour evidence source is missing: {source.name}")
    return source


def write_evidence(source: Path, destination: Path) -> str:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as opened:
        image = ImageOps.exif_transpose(opened).convert("RGB")
        image.thumbnail((EVIDENCE_MAX_EDGE, EVIDENCE_MAX_EDGE), Image.Resampling.LANCZOS)
        image.save(destination, format="WEBP", quality=EVIDENCE_QUALITY, method=6, exif=b"", icc_profile=b"")
    with Image.open(destination) as output:
        if max(output.size) > EVIDENCE_MAX_EDGE or output.format != "WEBP":
            raise ValueError(f"Invalid generated evidence image: {destination.name}")
    return sha256(destination)


def build_snapshot(workbook: Path, gps_inventory: Path, mission_root: Path, evidence_dir: Path) -> dict[str, object]:
    health_rows = workbook_rows(workbook)
    source_rows = inventory_rows(gps_inventory)
    captures = unique_green_captures(source_rows)
    observations: list[dict[str, object]] = []
    evidence_paths: set[Path] = set()

    for health, capture in zip(health_rows, captures, strict=True):
        area_index = int(capture["folder_index"])
        area_id = f"MPOC-SURVEY-{area_index:03d}"
        if area_index == 1:
            display_status = "Infected"
            status_source = "user-designated"
        elif health["Severity"] == "Severe":
            display_status = "Suspected"
            status_source = "modelled-risk"
        else:
            display_status = "Healthy"
            status_source = "modelled-risk"

        evidence_image = None
        evidence_hash = None
        if display_status in {"Infected", "Suspected"}:
            evidence_name = f"{health['Tree ID']}.webp"
            destination = evidence_dir / evidence_name
            source = evidence_source_for(capture, source_rows, mission_root)
            evidence_hash = write_evidence(source, destination)
            evidence_image = f"assets/mapped-poc/evidence/{evidence_name}"
            evidence_paths.add(destination.resolve())

        observations.append(
            {
                "treeId": health["Tree ID"],
                "areaId": area_id,
                "origin": "source-folder",
                "captureUuid": capture["capture_uuid"],
                "captureNumber": capture["capture_number"],
                "capturedAt": capture["capture_timestamp"],
                "latitude": round(float(capture["latitude"]), 9),
                "longitude": round(float(capture["longitude"]), 9),
                "ganodermaRiskScore": health["Ganoderma Risk Score"],
                "healthStatus": health["Health Status"],
                "severity": health["Severity"],
                "displayStatus": display_status,
                "statusSource": status_source,
                "evidenceImage": evidence_image,
                "evidenceSha256": evidence_hash,
            }
        )

    for index in range(113, 138):
        risk_score = float(66 + ((index * 7) % 30))
        observations.append(
            {
                "treeId": f"TREE-{index:04d}",
                "areaId": "MPOC-SURVEY-001",
                "origin": "layout-only",
                "captureUuid": None,
                "captureNumber": None,
                "capturedAt": None,
                "latitude": None,
                "longitude": None,
                "ganodermaRiskScore": risk_score,
                "healthStatus": "Unhealthy",
                "severity": "Severe",
                "displayStatus": "Infected",
                "statusSource": "deterministic-modelled",
                "evidenceImage": None,
                "evidenceSha256": None,
            }
        )

    areas = []
    for index in range(1, 6):
        mission = [row for row in captures if row["folder_index"] == index]
        area_id = f"MPOC-SURVEY-{index:03d}"
        area_observations = [row for row in observations if row["areaId"] == area_id]
        coordinate_points = [(float(row["latitude"]), float(row["longitude"])) for row in mission]
        areas.append(
            {
                "id": area_id,
                "name": f"Survey Area {index:03d}",
                "sourceMission": mission[0]["folder"],
                "centroid": {
                    "latitude": round(mean(point[0] for point in coordinate_points), 9),
                    "longitude": round(mean(point[1] for point in coordinate_points), 9),
                },
                "geofenceKind": "camera-footprint-convex-hull",
                "geofence": convex_hull(coordinate_points),
                "observationCount": len(area_observations),
                "observationIds": [row["treeId"] for row in area_observations],
            }
        )

    status_counts = {
        status: sum(row["displayStatus"] == status for row in observations)
        for status in ("Healthy", "Suspected", "Infected")
    }
    area_counts = [area["observationCount"] for area in areas]
    if len(observations) != 137 or area_counts != DISPLAY_AREA_COUNTS:
        raise ValueError(f"Expected 137 display observations split {DISPLAY_AREA_COUNTS}; received {area_counts}")
    if status_counts != {"Healthy": 65, "Suspected": 20, "Infected": 52}:
        raise ValueError(f"Unexpected display-status counts: {status_counts}")
    if len({row["treeId"] for row in observations}) != len(observations):
        raise ValueError("Tree IDs must be globally unique")
    if any(
        row["origin"] == "layout-only"
        and any(
            row[field] is not None
            for field in (
                "captureUuid",
                "capturedAt",
                "latitude",
                "longitude",
                "evidenceImage",
                "evidenceSha256",
            )
        )
        for row in observations
    ):
        raise ValueError("Layout-only observations must not contain source or evidence values")
    if any(
        (row["origin"] == "source-folder" and row["displayStatus"] in {"Infected", "Suspected"})
        != bool(row["evidenceImage"] and row["evidenceSha256"])
        for row in observations
    ):
        raise ValueError("Evidence must exist only for infected or suspected observations")
    generated_paths = {path.resolve() for path in evidence_dir.glob("*.webp")}
    if generated_paths != evidence_paths:
        raise ValueError("Evidence directory contains missing or stale WebP files")

    return {
        "version": 2,
        "provenance": {
            "captureDate": "2026-06-20",
            "sourceMissions": [area["sourceMission"] for area in areas],
            "workbookSha256": WORKBOOK_SHA256,
            "gpsInventorySha256": sha256(gps_inventory),
            "coordinateKind": "camera-exposure",
            "geofenceKind": "camera-footprint-convex-hull",
            "riskKind": "deterministic-modelled",
            "statusSemantics": "All Survey Area 001 observations are infected. Source observations are user-designated; layout-only observations use deterministic modelled risk and require browser-local operational positions. Severe modelled-risk observations in other areas are suspected.",
            "evidenceKind": "Natural-colour DJI D images resized to at most 1280 pixels and stored as repository WebP assets for infected and suspected observations only.",
            "runtimeDependency": "The snapshot and evidence assets are repository-owned and require no runtime access to external source drives.",
            "notice": "Camera exposure positions and derived display geofences are not surveyed palm-base coordinates or legal farm boundaries. Ganoderma scores are modelled POC values, not field or laboratory diagnoses.",
        },
        "areas": areas,
        "observations": observations,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--gps-inventory", required=True, type=Path)
    parser.add_argument("--mission-root", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--evidence-dir", type=Path)
    args = parser.parse_args()
    evidence_dir = args.evidence_dir or args.output.parent.parent / "assets" / "mapped-poc" / "evidence"
    snapshot = build_snapshot(args.workbook, args.gps_inventory, args.mission_root, evidence_dir)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(snapshot, ensure_ascii=False, separators=(",", ":"))
    args.output.write_text(
        "// Generated by scripts/build_mapped_poc_snapshot.py; do not edit by hand.\n"
        f"const MAPPED_POC_DATA = Object.freeze({payload});\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
